#!/usr/bin/env python3
"""Update the Bhikkhunī Pātimokkha Excel sheet (e.g. BiPv7.xlsx) from two Anki decks.

TARGET deck: "Pali Bhikkhunī Patimokkha Memorizing for Bhikkhunī"
SOURCE deck: "Pali Bhikkhunī Patimokkha Memorizing"

1. Rows whose noID is in the TARGET deck get Alt Pali / Alt English /
   Alt Audio / Pali / English / Audio refreshed from it. An empty
   Alt Pali field is filled with the Pali text.
2. SOURCE notes with no TARGET counterpart are inserted (noID kept as in
   the SOURCE deck, e.g. "BI VIN PAT SK 51-55"). A range "X a-b" goes
   right after the last single/pair/sub-rule row of section X whose end
   rule is <= b: SK 51-55 after SK 55, SK 1-10 after SK 9 10, AS 1-7
   after AS 3.
3. Rows found in neither deck are deleted.
4. "order" is renumbered 1..N.

Reads the decks live from Anki through the AnkiConnect add-on (Anki must
be open), or from Anki "Notes in Plain Text" exports with --txt.

Usage:
  python3 update_bip.py BiPv7.xlsx BiPv8.xlsx
  python3 update_bip.py BiPv7.xlsx BiPv8.xlsx --txt TARGET.txt SOURCE.txt

Needs: pip install openpyxl
"""
import argparse
import copy
import csv
import html
import json
import re
import sys
import unicodedata
import urllib.request

import openpyxl

TARGET_DECK = "Pali Bhikkhunī Patimokkha Memorizing for Bhikkhunī"
SOURCE_DECK = "Pali Bhikkhunī Patimokkha Memorizing"
ANKI_URL = "http://127.0.0.1:8765"

# Column numbers in an Anki text export (deck name is column 0, then the
# note fields in note-type order). With AnkiConnect, field order k is
# column k + 1.
F_ID, F_PALI, F_ALT_PALI, F_ENGLISH, F_AUDIO, F_ALT_AUDIO, F_ALT_ENGLISH = 1, 3, 4, 7, 10, 18, 19
COLUMNS = ["order", "noID", "Alt Pali", "Alt English", "Alt Audio", "Pali", "English", "Audio"]


# ── Reading the decks ─────────────────────────────────────────────────────────
def anki(action, **params):
    req = json.dumps({"action": action, "version": 6, "params": params}).encode("utf-8")
    try:
        with urllib.request.urlopen(urllib.request.Request(ANKI_URL, req), timeout=60) as resp:
            reply = json.load(resp)
    except OSError as e:
        sys.exit(f"Cannot reach AnkiConnect at {ANKI_URL} ({e}). Is Anki open with AnkiConnect installed?")
    if reply.get("error"):
        sys.exit(f"AnkiConnect {action}: {reply['error']}")
    return reply["result"]


def deck_from_anki(deck):
    deck = unicodedata.normalize("NFC", deck)
    ids = anki("findNotes", query=f'"deck:{deck}"')
    if not ids:
        sys.exit(f'No notes found in deck "{deck}" (check the exact deck name in Anki).')
    notes = {}
    for info in anki("notesInfo", notes=ids):
        fields = sorted(info["fields"].values(), key=lambda f: f["order"])
        row = [deck] + [f["value"] for f in fields]
        notes[row[F_ID]] = row
    return notes


def deck_from_txt(path):
    with open(path, encoding="utf-8") as f:
        lines = [l for l in f.read().split("\n") if l and not l.startswith("#")]
    return {r[F_ID]: r for r in csv.reader(lines, delimiter="\t", quotechar='"')}


# ── Conversions ───────────────────────────────────────────────────────────────
def to_cell(field):
    """Anki HTML -> Excel text, matching how the existing sheet was made."""
    s = re.sub(r"<br\s*/?>", "\n", field, flags=re.I)
    s = re.sub(r"<[^>]+>", "", s)          # <div>, <b>, <i>, ... are dropped
    return html.unescape(s)                 # &nbsp; -> non-breaking space, &amp; -> &


def note_values(note):
    pali = to_cell(note[F_PALI])
    return {
        "Alt Pali": to_cell(note[F_ALT_PALI]) or pali,   # empty Alt Pali -> copy Pali
        "Alt English": to_cell(note[F_ALT_ENGLISH]),
        "Alt Audio": note[F_ALT_AUDIO],
        "Pali": pali,
        "English": to_cell(note[F_ENGLISH]),
        "Audio": note[F_AUDIO],
    }


def source_to_target_id(sid):
    """'BI VIN PAT SK 1-10' -> 'BI VIN PAT BI SK 1-10'; 'BI VIN INTRO' -> 'BI VIN BI INTRO'."""
    if sid == "BI VIN INTRO":
        return "BI VIN BI INTRO"
    if sid.startswith("BI VIN PAT "):
        return "BI VIN PAT BI " + sid[len("BI VIN PAT "):]
    return sid


RANGE = re.compile(r"^(?P<pre>.* )(?P<sec>[A-Z]+) (?P<a>\d+)-(?P<b>\d+)$")
ATOMIC = re.compile(r"^(?P<pre>.* )(?P<sec>[A-Z]+) (?P<n>\d+)(?:\.\d+| (?P<m>\d+))?$")


def atomic_end(noid, prefix, sec):
    m = ATOMIC.match(noid)
    if not m or m["pre"] != prefix or m["sec"] != sec:
        return None
    return int(m["m"] or m["n"])


# ── Main ──────────────────────────────────────────────────────────────────────
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx_in")
    ap.add_argument("xlsx_out")
    ap.add_argument("--txt", nargs=2, metavar=("TARGET_TXT", "SOURCE_TXT"),
                    help="read Anki text exports instead of AnkiConnect")
    args = ap.parse_args()

    if args.txt:
        target, source = deck_from_txt(args.txt[0]), deck_from_txt(args.txt[1])
    else:
        target, source = deck_from_anki(TARGET_DECK), deck_from_anki(SOURCE_DECK)
    source_as_target = {source_to_target_id(sid) for sid in source}

    wb = openpyxl.load_workbook(args.xlsx_in)
    ws = wb.active
    header = [c.value for c in ws[1]]
    if header[:len(COLUMNS)] != COLUMNS:
        sys.exit(f"Unexpected header: {header}")
    col = {name: i for i, name in enumerate(COLUMNS)}
    old_count = ws.max_row - 1

    rows = [list(r) for r in ws.iter_rows(min_row=2, max_col=len(COLUMNS), values_only=True)]
    styles = [[copy.copy(c._style) for c in r[:len(COLUMNS)]] for r in ws.iter_rows(min_row=2)]

    # 1. Refresh rows from the TARGET deck; 3. drop rows in neither deck
    changed, deleted, kept_rows, kept_styles = 0, [], [], []
    for r, st in zip(rows, styles):
        noid = r[col["noID"]]
        note = target.get(noid)
        if note is None and noid not in source and noid not in source_as_target:
            deleted.append(noid)
            continue
        if note is not None:
            for name, value in note_values(note).items():
                if (r[col[name]] or "") != value:
                    r[col[name]] = value
                    changed += 1
        kept_rows.append(r)
        kept_styles.append(st)
    rows, styles = kept_rows, kept_styles

    # 2. Insert SOURCE notes missing from the TARGET deck
    present = {r[col["noID"]] for r in rows}
    inserts, unplaced = [], []
    for sid, note in source.items():
        tid = source_to_target_id(sid)
        if tid in target or sid in present or tid in present:
            continue
        m = RANGE.match(tid)
        anchor = None
        if m:
            for i, r in enumerate(rows):
                end = atomic_end(r[col["noID"]], m["pre"], m["sec"])
                if end is not None and end <= int(m["b"]):
                    anchor = i
        if anchor is None:
            unplaced.append(sid)
            continue
        new = [None] * len(COLUMNS)
        new[col["noID"]] = sid
        for name, value in note_values(note).items():
            new[col[name]] = value
        inserts.append((anchor, int(m["a"]), new))
    if unplaced:
        sys.exit(f"Could not place these notes (not 'X a-b' ranges with a matching row): {unplaced}")

    # Insert bottom-up so earlier anchors keep their positions; ranges sharing
    # an anchor end up in ascending start order.
    for anchor, _, new in sorted(inserts, key=lambda t: (t[0], t[1]), reverse=True):
        rows.insert(anchor + 1, new)
        styles.insert(anchor + 1, [copy.copy(s) for s in styles[anchor]])
        print(f"inserted {new[col['noID']]:24s} after {rows[anchor][col['noID']]}")

    # 4. Renumber and write back
    for i, r in enumerate(rows, start=1):
        r[col["order"]] = str(i)
    for i, (r, st) in enumerate(zip(rows, styles), start=2):
        for j, value in enumerate(r, start=1):
            cell = ws.cell(row=i, column=j, value=value)
            cell._style = copy.copy(st[j - 1])
    if old_count > len(rows):
        ws.delete_rows(len(rows) + 2, old_count - len(rows))
    wb.save(args.xlsx_out)

    print(f"\n{changed} cells updated from the target deck")
    print(f"{len(inserts)} rows inserted from the source deck")
    print(f"{len(deleted)} rows deleted (in neither deck): {deleted or 'none'}")
    print(f"Saved {args.xlsx_out} ({len(rows)} rows)")


if __name__ == "__main__":
    main()
